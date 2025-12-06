"""
ExcelGenerator for creating column-based Excel files.

This module provides the ExcelGenerator class that creates Excel workbooks
with a column-based format: one column per test parameter, one row per device.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from stdf_mcp.tools.filtered_data.models import ExcelWorkbook, ExcelWorksheet


class ExcelGenerator:
    """
    Generates column-based Excel files for filtered test data.

    Format:
    - Single worksheet: "Filtered_Test_Data"
    - Column headers: common columns (Device_ID, Site, Head, Timestamp) + test parameters
    - One row per device in chronological order
    - Sparse data: empty cells for tests not run on specific devices
    """

    def __init__(self) -> None:
        """Initialize Excel generator."""
        pass

    def generate_workbook(
        self, excel_data: ExcelWorkbook, output_path: Path
    ) -> Path:
        """
        Generate Excel workbook and save to file.

        Args:
            excel_data: ExcelWorkbook with worksheet and metadata
            output_path: Path where Excel file will be saved

        Returns:
            Path to created Excel file

        Raises:
            ValueError: If column count exceeds Excel limit
        """
        # Validate column limit (T030)
        self._validate_column_limit(excel_data.worksheet)

        # Create workbook
        wb = Workbook()

        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        # Create main data worksheet
        ws = wb.create_sheet("Filtered_Test_Data")
        self._write_worksheet(ws, excel_data.worksheet)

        # Save workbook
        wb.save(output_path)

        return output_path

    def _validate_column_limit(self, worksheet: ExcelWorksheet) -> None:
        """
        Validate that total columns don't exceed Excel limit (T030).

        Args:
            worksheet: ExcelWorksheet to validate

        Raises:
            ValueError: If column count exceeds 16,384
        """
        total_columns = worksheet.total_columns

        if total_columns > 16384:
            raise ValueError(
                f"TOO_MANY_COLUMNS: Column count {total_columns} exceeds Excel limit 16,384. "
                f"Common columns: {len(worksheet.common_columns)}, "
                f"Test parameters: {len(worksheet.test_parameters)}. "
                f"Reduce the number of filtered test parameters."
            )

    def _write_worksheet(
        self, ws: Worksheet, worksheet_data: ExcelWorksheet
    ) -> None:
        """
        Write column-based data to worksheet.

        Args:
            ws: openpyxl Worksheet object
            worksheet_data: ExcelWorksheet with data to write
        """
        # Write column headers (T029)
        headers = worksheet_data.get_column_headers()
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)

        # Get ordered parameter names for row generation
        ordered_params = worksheet_data.get_ordered_parameter_names()

        # Write device rows in chronological order (T028)
        for row_idx, device in enumerate(worksheet_data.device_rows, start=2):
            row_data = device.to_excel_row(ordered_params)

            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    def generate_with_metadata(
        self,
        excel_data: ExcelWorkbook,
        output_path: Path,
        include_metadata: bool = True,
    ) -> Path:
        """
        Generate Excel workbook with optional metadata worksheet.

        Args:
            excel_data: ExcelWorkbook with worksheet and metadata
            output_path: Path where Excel file will be saved
            include_metadata: Whether to include _Metadata worksheet

        Returns:
            Path to created Excel file
        """
        # Validate column limit
        self._validate_column_limit(excel_data.worksheet)

        # Create workbook
        wb = Workbook()

        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        # Create main data worksheet
        ws = wb.create_sheet("Filtered_Test_Data")
        self._write_worksheet(ws, excel_data.worksheet)

        # Create metadata worksheet if requested
        if include_metadata:
            self._write_metadata_worksheet(wb, excel_data)

        # Save workbook
        wb.save(output_path)

        return output_path

    def _write_metadata_worksheet(
        self, wb: Workbook, excel_data: ExcelWorkbook
    ) -> None:
        """
        Write metadata worksheet with filter criteria and statistics.

        Args:
            wb: openpyxl Workbook object
            excel_data: ExcelWorkbook with metadata
        """
        ws = wb.create_sheet("_Metadata")

        # Get metadata dictionary
        metadata = excel_data.generate_metadata_dict()

        # Write metadata sections
        row = 1
        for section_name, section_data in metadata.items():
            # Section header
            cell = ws.cell(row=row, column=1, value=section_name)
            cell.font = Font(bold=True, size=12)
            row += 1

            # Section content
            if isinstance(section_data, dict):
                for key, value in section_data.items():
                    ws.cell(row=row, column=1, value=key)
                    ws.cell(row=row, column=2, value=str(value))
                    row += 1
            else:
                ws.cell(row=row, column=1, value=str(section_data))
                row += 1

            # Blank row between sections
            row += 1
