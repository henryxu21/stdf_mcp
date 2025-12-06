"""Excel workbook generator for test time data.

This module handles generation of Excel (.xlsx) files from test time records,
following the column-based format used by the filtered_data tool.
"""

from datetime import datetime
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from stdf_mcp.tools.test_time.models import TestTimeRecord, TestTimeWorkbook


class ExcelTestTimeGenerator:
    """Generates Excel workbook from test time data.

    Creates .xlsx files with columns: Device_ID, Site, Head, Timestamp, Test_Time
    following the existing filtered_data Excel format. Empty cells used for
    missing test times (None values).
    """

    # Column headers per FR-003 specification
    COLUMN_HEADERS = ["Device_ID", "Site", "Head", "Timestamp", "Test_Time"]

    def generate_excel(self, workbook: TestTimeWorkbook) -> Path:
        """Generate Excel file from test time workbook data.

        Creates .xlsx file in output_path location with test time data
        in column-based format.

        Args:
            workbook: TestTimeWorkbook containing records and metadata

        Returns:
            Path to generated Excel file

        Raises:
            OSError: If file cannot be written
        """
        # Create new workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Time Data"

        # Write column headers
        self._write_headers(ws)

        # Write data rows
        self._write_rows(ws, workbook.records)

        # Save workbook
        try:
            wb.save(workbook.output_path)
        except OSError as e:
            raise OSError(
                f"Failed to write Excel file to {workbook.output_path}: {e}"
            )

        return workbook.output_path

    def _write_headers(self, ws: Worksheet) -> None:
        """Write column headers to first row.

        Args:
            ws: Worksheet to write headers to
        """
        for col_idx, header in enumerate(self.COLUMN_HEADERS, start=1):
            ws.cell(row=1, column=col_idx, value=header)

    def _write_rows(self, ws: Worksheet, records: List[TestTimeRecord]) -> None:
        """Write test time records as data rows.

        Each record becomes one row with columns matching header order.
        Missing test times (None) written as empty cells.

        Args:
            ws: Worksheet to write rows to
            records: List of TestTimeRecord objects in chronological order
        """
        for row_idx, record in enumerate(records, start=2):  # Start after header
            # Column 1: Device_ID
            ws.cell(row=row_idx, column=1, value=record.device_id)

            # Column 2: Site
            ws.cell(row=row_idx, column=2, value=record.site)

            # Column 3: Head
            ws.cell(row=row_idx, column=3, value=record.head)

            # Column 4: Timestamp (always empty - PIR lacks timestamps)
            ws.cell(row=row_idx, column=4, value=None)

            # Column 5: Test_Time (None creates empty cell per FR-006)
            ws.cell(row=row_idx, column=5, value=record.test_time_seconds)

    @staticmethod
    def generate_filename(stdf_path: Path, timestamp: datetime) -> str:
        """Generate Excel filename from STDF file and timestamp.

        Follows pattern: <stdf_filename>_TestTime_<YYYYMMDD_HHMMSS>.xlsx

        Args:
            stdf_path: Path to source STDF file
            timestamp: Timestamp for filename

        Returns:
            Generated filename (not full path, just filename)

        Example:
            >>> generate_filename(Path("example_1.stdf"), datetime(2025, 1, 27, 14, 35, 22))
            "example_1_TestTime_20250127_143522.xlsx"
        """
        stdf_name = stdf_path.stem  # Filename without extension
        timestamp_str = timestamp.strftime("%Y%m%d_%H%M%S")
        return f"{stdf_name}_TestTime_{timestamp_str}.xlsx"

    @staticmethod
    def create_output_path(
        stdf_file_path: Path,
        output_directory: Path,
        timestamp: datetime
    ) -> Path:
        """Create full output path for Excel file.

        Args:
            stdf_file_path: Source STDF file path
            output_directory: Directory for Excel output
            timestamp: Timestamp for filename

        Returns:
            Full path to Excel output file
        """
        filename = ExcelTestTimeGenerator.generate_filename(
            stdf_file_path, timestamp
        )
        return output_directory / filename
